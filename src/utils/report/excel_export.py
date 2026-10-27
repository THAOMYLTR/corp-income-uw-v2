"""Write the result as an Excel workbook laid out like the BRD.

Sheets: Tóm tắt · 1. Pháp lý cơ bản · 2. Cổ đông · 3. Người đại diện ·
4. Nợ thuế - tranh chấp · 5. Tài chính (BCTC, chỉ số, cảnh báo, thu nhập ghi nhận) ·
Mô hình HĐKD [BRD 27/10] · Kiểm tra (all findings) · Hồ sơ (files, how each was read, its
document-matrix code, agents and validation rules). Amounts are numbers in đồng with a thousands format, so they can be
summed and filtered.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.agents.calculator.legal_checks import SOURCE_LABELS

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
SEVERITY_FILL = {
    "alert": PatternFill("solid", fgColor="F8CBAD"),
    "warning": PatternFill("solid", fgColor="FFE699"),
    "info": PatternFill("solid", fgColor="DDEBF7"),
    "ok": PatternFill("solid", fgColor="E2EFDA"),
}
SEVERITY_LABELS = {"alert": "Cảnh báo", "warning": "Cần làm rõ", "info": "Ghi nhận", "ok": "Khớp"}
THIN = Side(style="thin", color="A6A6A6")
MONEY_FORMAT = '#,##0;(#,##0);"-"'
MONEY_KEYS = {"amount", "claim_amount", "capital_amount"}


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]], widths: list[int],
           money_cols: set[int] = frozenset()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append([("" if v is None else v) for v in row])
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if cell.column in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = MONEY_FORMAT
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return ws


def write_excel(result: dict[str, Any], path: str | Path) -> Path:
    items = {int(k): v for k, v in result["items"].items()}
    wb = Workbook()
    wb.remove(wb.active)

    meta = result["meta"]
    _sheet(wb, "Tóm tắt", ["Mục", "Nội dung", "Kết luận"],
           [[n, item["title"], item["conclusion"]] for n, item in items.items()]
           + [[], ["", "Ngày tham chiếu", meta["as_of"]],
              ["", "KH cá nhân vay vốn", ", ".join(meta["borrower_names"]) or "—"]]
           # [BRD 16/10] groups (POC)
           + ([[], ["", "Phân nhóm doanh nghiệp", result["corp_group"]["group"]],
               ["", "Phân nhóm tài chính", result["corp_group"]["finance"]],
               ["", "Phân nhóm bureau", result["corp_group"]["bureau"]]]
              + [["", "Tiêu chí Bad corp", r] for r in result["corp_group"]["bad_reasons"]]
              if result.get("corp_group") else [])
           # [BRD 16/10] overall credit view
           + ([[], ["", "Đề xuất", (result.get("credit_assessment") or {}).get("quyet_dinh", "")],
               ["", "Tóm tắt", (result.get("credit_assessment") or {}).get("tom_tat", "")]]
              + [["", f"Rủi ro ({r.get('muc_do', '')})", f"{r.get('noi_dung', '')} — {r.get('can_cu', '')}"]
                 for r in (result.get("credit_assessment") or {}).get("rui_ro", []) if isinstance(r, dict)]
              + [["", "Điều kiện", d] for d in (result.get("credit_assessment") or {}).get("dieu_kien", [])]
              if result.get("credit_assessment") else []),
           [6, 45, 100])

    _sheet(wb, "1. Pháp lý cơ bản", ["Thông tin", "Giá trị", "Nguồn"],
           [[label, entry["value"], entry["source"]] for label, entry in items[1]["summary"].items()],
           [32, 55, 55])

    members = items[2]["summary"].get("members", [])
    _sheet(wb, "2. Cổ đông",
           ["Thành viên/cổ đông", "Loại", "CCCD/MSDN", "Vốn góp (đồng)", "Tỷ lệ (%)", "Tỷ lệ là",
            "Ngày nhận vốn góp", "Hình thức", "Nguồn"],
           [[m["name"], "Pháp nhân" if m["member_type"] == "phap_nhan" else "Cá nhân", m["id_number"],
             m["capital_amount"], m["ownership_pct"], m["pct_source"], m["acquisition_date"],
             m["acquisition_method"], m["source"]] for m in members],
           [30, 11, 16, 18, 10, 11, 16, 20, 40], money_cols={4})

    summary3 = items[3]["summary"]
    rows3 = [[r.get("name"), r.get("title"), r.get("id_number"), r.get("appointed_date"),
              "Người đại diện theo pháp luật", r.get("source")]
             for r in summary3.get("legal_representatives", [])]
    if summary3.get("actual_operator"):
        op = summary3["actual_operator"]
        rows3.append([op.get("name"), op.get("role"), "", "", "Người điều hành thực tế", op.get("source")])
    _sheet(wb, "3. Người đại diện", ["Họ tên", "Chức danh", "CCCD", "Ngày bổ nhiệm", "Vai trò", "Nguồn"],
           rows3, [28, 22, 16, 16, 28, 45])

    summary4 = items[4]["summary"]
    groups = (("Nợ thuế", "tax_debts"), ("Nợ BHXH", "social_insurance_debts"),
              ("Xác nhận không nợ", "clearance_confirmations"), ("Xử phạt", "administrative_penalties"),
              ("Cưỡng chế", "enforcement_actions"), ("Kiểm tra/thanh tra", "inspections"),
              ("Tranh chấp", "disputes"), ("Rủi ro tại thực địa", "site_risks"))
    rows4 = []
    for label, key in groups:
        for row in summary4.get(key, []):
            amount = next((row.get(k) for k in ("amount", "claim_amount") if row.get(k) is not None), None)
            detail = "; ".join(f"{k}: {v}" for k, v in row.items()
                               if k not in ("source", "amount", "claim_amount", "page") and v not in (None, "", []))
            rows4.append([label, amount, detail, row.get("page"), row.get("source")])
    _sheet(wb, "4. Nợ thuế - tranh chấp", ["Loại", "Số tiền (đồng)", "Chi tiết", "Trang", "Nguồn"],
           rows4, [20, 18, 90, 7, 40], money_cols={2})

    fin = result.get("financials") or {}
    if fin:
        periods = {**fin.get("years", {}), **({fin["interim"]["period"]: fin["interim"]} if fin.get("interim") else {})}
        labels = (("Tổng tài sản", "total_assets"), ("Tài sản ngắn hạn", "current_assets"),
                  ("Tài sản dài hạn", "non_current_assets"), ("Nợ phải trả", "total_liabilities"),
                  ("Nợ ngắn hạn", "current_liabilities"), ("Vốn chủ sở hữu", "equity"),
                  ("Tổng nguồn vốn", "total_capital"), ("Doanh thu thuần", "net_revenue"),
                  ("Giá vốn hàng bán", "cogs"), ("Lợi nhuận gộp", "gross_profit"),
                  ("Doanh thu tài chính", "financial_income"), ("Chi phí tài chính", "financial_expenses"),
                  ("Chi phí bán hàng", "selling_expenses"), ("Chi phí quản lý DN", "admin_expenses"),
                  ("Lợi nhuận thuần", "operating_profit"), ("Lợi nhuận trước thuế", "profit_before_tax"),
                  ("Lợi nhuận sau thuế", "profit_after_tax"), ("Khấu hao", "depreciation"))
        rows5 = [[label] + [p.get(key) for p in periods.values()] for label, key in labels]
        for label, key in (("Biến động doanh thu", "revenue_growth"), ("LNST / DTT (ROS)", "ros"),
                           ("Hệ số thanh toán hiện hành", "current_ratio")):
            rows5.append([label] + [(fin.get("ratios", {}).get(y) or {}).get(key) for y in periods])
        rows5 += [[], ["Nguồn"] + ["; ".join(fin.get("sources") or [])]]
        rows5 += [["Cảnh báo", w] for w in fin.get("warnings") or []]
        income = result.get("income_estimate") or {}
        if income:
            proj = income.get("projection") or {}
            rows5 += [[], ["Cách tính", "Dự phóng theo PAKD" if proj else "Số liệu lịch sử"],
                      ["LNST hợp lý (đồng)", income.get("reasonable_lnst")]]
            if proj:   # [BRD 27/10]
                rows5 += [["Ngành TCB", proj.get("industry")],
                          ["DT dự phóng khả thi (đồng)", proj.get("feasible_revenue")],
                          ["ROS dự phóng khả thi", proj.get("feasible_ros")]]
            rows5 += [["Thu nhập DN ghi nhận (đồng)", income.get("value")],
                      ["Thành phần còn thiếu", "; ".join(income.get("missing") or []) or "—"]]
            rows5 += [["Cảnh báo", w] for w in income.get("warnings") or []]
        _sheet(wb, "5. Tài chính", ["Chỉ tiêu (đồng)"] + list(periods), rows5,
               [32] + [20] * max(1, len(periods)), money_cols=set(range(2, len(periods) + 2)))

    model = result.get("business_model") or {}
    if model:   # [BRD 27/10] mục 2 — partners of the diagram, largest first
        rows_m = [["Ngành TCB", (model.get("industry") or {}).get("name") or "Chưa xác định", None, None, None, ""]]
        for side, label in (("suppliers", "Đầu vào (nhà cung cấp)"), ("customers", "Đầu ra (khách mua)")):
            rows_m += [[label, p["name"], p["product"], p["value"], p["share_pct"], p.get("source")]
                       for p in model.get(side) or []]
        _sheet(wb, "Mô hình HĐKD", ["Bên", "Đối tác", "Mặt hàng", "Giá trị (đồng)", "Tỷ trọng (%)", "Nguồn"],
               rows_m, [24, 40, 26, 18, 12, 40], money_cols={4})

    findings = [[n, SEVERITY_LABELS[f["severity"]], f["message"], "\n".join(f.get("evidence") or [])]
                for n, item in items.items() for f in item["findings"]]
    ws = _sheet(wb, "Kiểm tra", ["Mục", "Mức", "Kết quả", "Căn cứ"], findings, [6, 12, 70, 90])
    severity_by_label = {v: k for k, v in SEVERITY_LABELS.items()}
    for row in ws.iter_rows(min_row=2):
        row[1].fill = SEVERITY_FILL[severity_by_label[row[1].value]]

    _sheet(wb, "Hồ sơ", ["File", "Thư mục", "Mã (ma trận)", "Loại", "Agent (ma trận)", "Nguồn", "Đọc bằng",
                         "Số ký tự", "Mục đã trích xuất", "Kiểm tra theo ma trận (người dùng xác nhận)", "Lỗi"],
           [[d["filename"], d.get("box") or "—", d.get("matrix_code") or "", d["label"],
             ", ".join(d.get("matrix_agents") or []), SOURCE_LABELS.get(d["source"], d["source"]),
             d.get("read_method", ""), d.get("chars", 0), ", ".join(d["passes_done"]),
             "\n".join(d.get("matrix_validation") or []),
             "; ".join(f"{k}: {v}" for k, v in d["errors"].items())] for d in result["documents"]],
           [34, 24, 10, 40, 22, 20, 16, 10, 34, 50, 50])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
