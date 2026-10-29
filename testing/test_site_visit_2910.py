"""[Sửa 29/10] Báo cáo khảo sát thực địa (mẫu MB, file Excel): KH vay, sản phẩm/dịch vụ."""

import os
import tempfile
import unittest
from types import SimpleNamespace

from openpyxl import Workbook

from src.agents.calculator import borrowers, business_model
from src.utils.reading.extractors import read_sheets


def _site_report(path: str) -> None:
    """The layout of the real report: labels stacked in one merged cell (B:F), values in G:K."""

    wb = Workbook()
    ws = wb.active
    ws.title = "MB - Báo cáo khảo sát thực địa"
    ws["B2"] = "Người đại diện theo PL"
    ws["G2"] = "Vũ Việt Dũng"
    ws["B4"] = "Thời gian khảo sát:\nĐịa điểm khảo sát:\nThành phần:\nĐại diện KH\nĐại diện TCB"
    ws["G4"] = "15/08/2026\nSố 324 Lê Duẩn, Đống Đa, Hà Nội\nSRBO + KH\nNguyễn Huệ Chi\nNGUYỄN QUANG HUY"
    ws.merge_cells("B4:F4")
    ws.merge_cells("G4:K4")
    ws.append([])
    ws.append(["", "Họ và tên", "Mối quan hệ", "Tỷ lệ góp vốn", "Ghi chú"])
    ws.append(["", "Vũ Hoàng Nguyên", "Chồng", "72.5%",
               "Ông Nguyên nắm vai trò quyết định chính.\nBác sỹ chuyên khoa ung bướu"])
    ws.append(["", "Sản phẩm/ dịch vụ", "Dịch vụ Y tế và hoạt động hỗ trợ y tế"])
    wb.save(path)


class SiteVisitTest(unittest.TestCase):
    def setUp(self):
        folder = tempfile.mkdtemp()
        self.path = os.path.join(folder, "bao_cao_thuc_dia.xlsx")
        _site_report(self.path)
        self.text = read_sheets(self.path)[0].content

    def test_label_and_value_on_one_row(self):
        row = next(line for line in self.text.splitlines() if "Đại diện KH" in line)
        self.assertIn("Nguyễn Huệ Chi", row)
        self.assertNotIn("NGUYỄN QUANG HUY", row)

    def test_note_with_line_break_stays_on_its_row(self):
        row = next(line for line in self.text.splitlines() if "Vũ Hoàng Nguyên" in line)
        self.assertIn("Bác sỹ chuyên khoa ung bướu", row)

    def test_borrower_is_the_customer_representative(self):
        doc = SimpleNamespace(source="field", content=self.text, filename="bao_cao_thuc_dia.xlsx",
                              declared_group="ho_so_noi_bo")
        members = [{"name": n, "member_type": "ca_nhan"}
                   for n in ("Vũ Hoàng Nguyên", "Nguyễn Huệ Chi", "Vũ Hoàng Tùng")]
        records = [{"pass": "shareholders", "source": "customer", "data": {"members": members}}]
        names, _ = borrowers.detect_borrowers([doc], records)
        self.assertEqual(names, ["Nguyễn Huệ Chi"])

    def test_borrower_without_registration_data(self):
        doc = SimpleNamespace(source="field", content=self.text, filename="x.xlsx", declared_group="")
        self.assertEqual(borrowers.detect_borrowers([doc], [])[0], ["Nguyễn Huệ Chi"])

    def test_bank_officer_is_never_taken(self):
        line = "Đại diện TCB\tĐại diện TCB\tNGUYỄN QUANG HUY"
        self.assertEqual(borrowers._name_after_label(line, {}), "")

    def test_product_service_from_site_report(self):
        records = [{"pass": "business_model", "document_type": "hop_dong_dau_vao_dau_ra",
                    "data": {"product_service": ""}},
                   {"pass": "business_model", "document_type": "bao_cao_thuc_dia",
                    "data": {"product_service": "Dịch vụ Y tế và hoạt động hỗ trợ y tế"}}]
        self.assertEqual(business_model.product_service(records), "Dịch vụ Y tế và hoạt động hỗ trợ y tế")

    def test_share_is_value_over_total_of_identified_partners(self):
        records = [{"pass": "business_model", "document_type": "hop_dong_dau_vao_dau_ra", "data": {
            "suppliers": [{"name": "Dược Hậu Giang", "value": 135e6, "share_pct": 90},
                          {"name": "Nam Phùng", "value": 45e6},
                          {"name": "Không rõ giá trị", "value": None}]}}]
        rows = {p["name"]: p["share_pct"] for p in business_model.partners(records, "suppliers")}
        self.assertEqual(rows, {"Dược Hậu Giang": 75.0, "Nam Phùng": 25.0, "Không rõ giá trị": None})


if __name__ == "__main__":
    unittest.main()
